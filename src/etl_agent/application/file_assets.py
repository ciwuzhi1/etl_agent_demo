"""文件资产服务：文件名净化、格式白名单、路径生成与状态机。

课件 §5.4.1（第 5 章）要求把客户端文件名只当"显示别名"：先砍掉路径段和控制字符，
再由资产 UUID + 受信扩展名生成 `<uuid>/source.csv`；未知扩展名立即失败，绝不"当文本试试看"。
§5.6 要求 FileAssetService 用 SELECT FOR UPDATE 串行化上传/解析/完成/失败的状态变化。
"""

import csv
import io
import json
import re
import uuid
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from pathlib import PurePosixPath
from typing import Any, cast

import pyarrow as pa
import pyarrow.parquet as arrow_parquet
import pyorc
from defusedxml import ElementTree as SafeElementTree
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from etl_agent.application.masking import SENSITIVE_COLUMN, mask_value
from etl_agent.application.profiling import require_project_role
from etl_agent.domain.enums import DataProfileStatus, FileAssetFormat, ProjectRole
from etl_agent.domain.errors import DomainError, ResourceNotFound
from etl_agent.infrastructure.models import AuditEvent, FileAsset

_EXTENSION_FORMATS = {
    ".csv": FileAssetFormat.CSV,
    ".json": FileAssetFormat.JSON,
    ".jsonl": FileAssetFormat.JSONL,
    ".ndjson": FileAssetFormat.JSONL,
    ".xlsx": FileAssetFormat.XLSX,
    ".parquet": FileAssetFormat.PARQUET,
    ".orc": FileAssetFormat.ORC,
    ".xml": FileAssetFormat.XML,
    ".md": FileAssetFormat.MARKDOWN,
    ".markdown": FileAssetFormat.MARKDOWN,
    ".pdf": FileAssetFormat.PDF,
}


@dataclass(frozen=True)
class ParsedFileProfile:
    """保存文件解析后的统一行模型、Schema 与非敏感统计。"""

    rows: list[dict[str, Any]]
    schema_definition: dict[str, Any]
    statistics: dict[str, Any]
    masked_sample: list[dict[str, Any]]


class FileAssetService:
    """编排文件资产登记、异步解析状态与项目权限边界。"""

    async def request(
        self,
        session: AsyncSession,
        *,
        project_id: uuid.UUID,
        actor_id: uuid.UUID,
        filename: str,
        media_type: str | None,
        trace_id: str | None = None,
    ) -> FileAsset:
        """登记待上传资产；此阶段只冻结安全文件名和格式，不接触内容。"""
        await require_project_role(
            session,
            project_id,
            actor_id,
            {ProjectRole.ADMIN, ProjectRole.DATA_ENGINEER},
        )
        safe_filename = _safe_filename(filename)
        file_format = format_for_filename(safe_filename)
        asset = FileAsset(
            project_id=project_id,
            original_filename=safe_filename,
            file_format=file_format,
            media_type=(media_type or "application/octet-stream")[:255],
            created_by=actor_id,
        )
        session.add(asset)
        await session.flush()
        session.add(
            AuditEvent(
                actor_id=actor_id,
                project_id=project_id,
                action="file_asset.requested",
                resource_type="FileAsset",
                resource_id=asset.id,
                before=None,
                after={"filename": safe_filename, "format": file_format.value},
                trace_id=trace_id,
            )
        )
        await session.flush()
        return asset

    async def record_stored(
        self,
        session: AsyncSession,
        *,
        asset_id: uuid.UUID,
        source_uri: str,
        byte_size: int,
        sha256: str,
    ) -> FileAsset:
        """在上传原子落盘后记录摘要；路径只能来自受管存储适配器。"""
        asset = await self._get_for_update(session, asset_id)
        if asset.status is not DataProfileStatus.PENDING or asset.source_uri is not None:
            raise DomainError("文件资产当前状态不能登记上传结果")
        asset.source_uri = source_uri
        asset.byte_size = byte_size
        asset.sha256 = sha256
        await session.flush()
        await session.refresh(asset)
        return asset

    async def mark_running(self, session: AsyncSession, asset_id: uuid.UUID) -> FileAsset:
        """锁定待处理资产并进入解析中状态，重复任务不会并发解析同一文件。"""
        asset = await self._get_for_update(session, asset_id)
        if asset.status is not DataProfileStatus.PENDING or not asset.source_uri:
            raise DomainError("只有上传完成的待处理文件可以开始解析")
        asset.status = DataProfileStatus.RUNNING
        asset.error = None
        await session.flush()
        return asset

    async def complete(
        self,
        session: AsyncSession,
        *,
        asset_id: uuid.UUID,
        normalized_uri: str,
        engine_path: str,
        profile: ParsedFileProfile,
        ttl_hours: int = 24,
    ) -> FileAsset:
        """提交规范化文件与脱敏 Profile，使其成为可选 Agent Source。"""
        asset = await self._get_for_update(session, asset_id)
        if asset.status is not DataProfileStatus.RUNNING:
            raise DomainError("文件资产当前状态不能写入解析结果")
        asset.normalized_uri = normalized_uri
        asset.engine_path = engine_path
        asset.status = DataProfileStatus.READY
        asset.schema_json = profile.schema_definition
        asset.statistics = profile.statistics
        asset.masked_sample = profile.masked_sample
        asset.row_count = len(profile.rows)
        asset.expires_at = datetime.now(UTC) + timedelta(hours=ttl_hours)
        asset.error = None
        session.add(
            AuditEvent(
                actor_id=asset.created_by,
                project_id=asset.project_id,
                action="file_asset.profiled",
                resource_type="FileAsset",
                resource_id=asset.id,
                before={"status": DataProfileStatus.RUNNING.value},
                after={
                    "status": DataProfileStatus.READY.value,
                    "format": asset.file_format.value,
                    "byte_size": asset.byte_size,
                    "row_count": asset.row_count,
                    "sha256": asset.sha256,
                },
            )
        )
        await session.flush()
        return asset

    async def fail(
        self,
        session: AsyncSession,
        *,
        asset_id: uuid.UUID,
        error: str,
    ) -> FileAsset:
        """记录不包含文件内容和主机路径的失败原因，并关闭本次解析。"""
        asset = await self._get_for_update(session, asset_id)
        asset.status = DataProfileStatus.FAILED
        asset.error = error[:2_000]
        await session.flush()
        return asset

    async def list_for_project(
        self,
        session: AsyncSession,
        *,
        project_id: uuid.UUID,
        actor_id: uuid.UUID,
    ) -> list[FileAsset]:
        """按项目权限列出文件资产，响应模型会排除内部 URI 与引擎路径。"""
        await require_project_role(
            session,
            project_id,
            actor_id,
            {
                ProjectRole.ADMIN,
                ProjectRole.DATA_ENGINEER,
                ProjectRole.APPROVER,
                ProjectRole.VIEWER,
            },
        )
        return list(
            (
                await session.scalars(
                    select(FileAsset)
                    .where(FileAsset.project_id == project_id)
                    .order_by(FileAsset.created_at.desc())
                    .limit(100)
                )
            ).all()
        )

    async def get_visible(
        self,
        session: AsyncSession,
        *,
        asset_id: uuid.UUID,
        actor_id: uuid.UUID,
    ) -> FileAsset:
        """读取单个文件资产并在返回前验证项目成员权限。"""
        asset = await session.get(FileAsset, asset_id)
        if asset is None:
            raise ResourceNotFound("文件资产不存在")
        await require_project_role(
            session,
            asset.project_id,
            actor_id,
            {
                ProjectRole.ADMIN,
                ProjectRole.DATA_ENGINEER,
                ProjectRole.APPROVER,
                ProjectRole.VIEWER,
            },
        )
        return asset

    @staticmethod
    async def _get_for_update(session: AsyncSession, asset_id: uuid.UUID) -> FileAsset:
        """使用行锁串行化上传、解析及重复任务状态变化。"""
        asset = (
            await session.execute(
                select(FileAsset).where(FileAsset.id == asset_id).with_for_update()
            )
        ).scalar_one_or_none()
        if asset is None:
            raise ResourceNotFound("文件资产不存在")
        return asset


def format_for_filename(filename: str) -> FileAssetFormat:
    """根据经过净化的扩展名选择固定解析器，拒绝双重猜测内容格式。"""
    suffix = PurePosixPath(filename).suffix.casefold()
    file_format = _EXTENSION_FORMATS.get(suffix)
    if file_format is None:
        supported = ", ".join(sorted(_EXTENSION_FORMATS))
        raise DomainError(f"不支持的文件格式；当前支持 {supported}")
    return file_format


def source_relative_path(asset_id: uuid.UUID, file_format: FileAssetFormat) -> str:
    """为原始上传生成只含资产 UUID 和受信扩展名的稳定路径。"""
    suffix = {
        FileAssetFormat.CSV: "csv",
        FileAssetFormat.JSON: "json",
        FileAssetFormat.JSONL: "jsonl",
        FileAssetFormat.XLSX: "xlsx",
        FileAssetFormat.PARQUET: "parquet",
        FileAssetFormat.ORC: "orc",
        FileAssetFormat.XML: "xml",
        FileAssetFormat.MARKDOWN: "md",
        FileAssetFormat.PDF: "pdf",
    }[file_format]
    return f"{asset_id}/source.{suffix}"


def engine_path_for(asset_id: uuid.UUID, engine_root: str) -> str:
    """生成 SeaTunnel 容器内只读挂载路径，禁止把宿主机路径交给模型。"""
    root = PurePosixPath(engine_root)
    if not root.is_absolute() or ".." in root.parts:
        raise DomainError("文件资产引擎根路径必须是安全的绝对路径")
    return str(root / str(asset_id) / "normalized.jsonl")


def _safe_filename(filename: str) -> str:
    """移除客户端路径片段并拒绝控制字符，防止文件名注入日志或审计。"""
    name = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not name or len(name) > 255 or any(ord(char) < 32 for char in name):
        raise DomainError("文件名为空、过长或包含控制字符")
    return name


_UNSAFE_COLUMN = re.compile(r"[^A-Za-z0-9_]+")


def parse_file_content(
    content: bytes,
    *,
    file_format: FileAssetFormat,
    max_rows: int,
    sample_limit: int = 20,
) -> ParsedFileProfile:
    """把异构文件解析为统一 JSON 行模型，并对样本实施字段级脱敏。"""
    rows, parser_statistics = _read_rows(content, file_format, max_rows)
    if not rows:
        raise DomainError("文件没有可解析的数据行")
    if len(rows) > max_rows:
        raise DomainError(f"文件超过单次解析允许的 {max_rows} 行上限")
    normalized_rows, mapping = _normalize_rows(rows)
    schema = _infer_schema(normalized_rows)
    masked_rows = [
        {mapping[str(key)]: mask_value(str(key), _json_value(value)) for key, value in row.items()}
        for row in rows[:sample_limit]
    ]
    return ParsedFileProfile(
        rows=normalized_rows,
        schema_definition={"columns": schema},
        statistics={
            "format": file_format.value,
            "row_count": len(normalized_rows),
            "column_count": len(schema),
            "column_mapping": mapping,
            **parser_statistics,
        },
        masked_sample=masked_rows,
    )


def normalized_relative_path(asset_id: uuid.UUID) -> str:
    """为规范化 JSONL 生成不含原始文件名的稳定受管路径。"""
    return f"{asset_id}/normalized.jsonl"


def _read_rows(
    content: bytes,
    file_format: FileAssetFormat,
    max_rows: int,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """路由到固定格式解析器，未知格式不会回退到宽松文本猜测。"""
    if file_format is FileAssetFormat.CSV:
        return _read_csv(content, max_rows), {}
    if file_format is FileAssetFormat.JSON:
        return _read_json(content, max_rows), {}
    if file_format is FileAssetFormat.JSONL:
        return _read_jsonl(content, max_rows), {}
    if file_format is FileAssetFormat.XLSX:
        return _read_xlsx(content, max_rows)
    if file_format is FileAssetFormat.PARQUET:
        return _read_parquet(content, max_rows)
    if file_format is FileAssetFormat.ORC:
        return _read_orc(content, max_rows)
    if file_format is FileAssetFormat.XML:
        return _read_xml(content, max_rows)
    if file_format is FileAssetFormat.MARKDOWN:
        return _read_markdown(content, max_rows), {"document_model": "markdown_elements"}
    if file_format is FileAssetFormat.PDF:
        return _read_pdf(content, max_rows)
    raise DomainError("没有可用的文件解析器")


def _decode_text(content: bytes) -> str:
    """只接受 UTF-8/UTF-8 BOM 文本，避免系统默认编码造成跨环境差异。"""
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DomainError("文本文件必须使用 UTF-8 编码") from exc


def _read_csv(content: bytes, max_rows: int) -> list[dict[str, Any]]:
    """解析带表头 CSV，并使用 Sniffer 识别常见分隔符。"""
    text = _decode_text(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames or any(not str(name or "").strip() for name in reader.fieldnames):
        raise DomainError("CSV 必须包含非空表头")
    rows: list[dict[str, Any]] = []
    for row in reader:
        rows.append(
            {
                str(key): (
                    value if SENSITIVE_COLUMN.search(str(key)) else _parse_text_scalar(value)
                )
                for key, value in row.items()
            }
        )
        _require_row_capacity(rows, max_rows)
    return rows


def _read_json(content: bytes, max_rows: int) -> list[dict[str, Any]]:
    """解析 JSON 对象或对象数组，拒绝无法形成字段行的标量数组。"""
    try:
        payload = json.loads(_decode_text(content))
    except json.JSONDecodeError as exc:
        raise DomainError("JSON 语法无效") from exc
    candidates = payload if isinstance(payload, list) else [payload]
    _require_row_capacity(candidates, max_rows)
    if not all(isinstance(item, dict) for item in candidates):
        raise DomainError("JSON 顶层必须是对象或对象数组")
    return [dict(item) for item in candidates]


def _read_jsonl(content: bytes, max_rows: int) -> list[dict[str, Any]]:
    """逐行解析 JSONL，每个非空行都必须是对象。"""
    rows: list[dict[str, Any]] = []
    for line_number, line in enumerate(_decode_text(content).splitlines(), start=1):
        if not line.strip():
            continue
        try:
            item = json.loads(line)
        except json.JSONDecodeError as exc:
            raise DomainError(f"JSONL 第 {line_number} 行语法无效") from exc
        if not isinstance(item, dict):
            raise DomainError(f"JSONL 第 {line_number} 行必须是对象")
        rows.append(dict(item))
        _require_row_capacity(rows, max_rows)
    return rows


def _read_xlsx(content: bytes, max_rows: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """以只读模式解析 XLSX 的首个工作表，并把首行作为字段名。"""
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise DomainError("XLSX 文件无法打开或结构已损坏") from exc
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        header = next(values, None)
        if header is None or not any(value is not None for value in header):
            return [], {"sheet_name": worksheet.title}
        names = [str(value or "").strip() for value in header]
        if any(not name for name in names):
            raise DomainError("XLSX 首行必须包含非空表头")
        rows: list[dict[str, Any]] = []
        for row in values:
            if not any(value is not None for value in row):
                continue
            rows.append(dict(zip(names, row, strict=False)))
            _require_row_capacity(rows, max_rows)
        return rows, {"sheet_name": worksheet.title}
    finally:
        workbook.close()


def _read_parquet(content: bytes, max_rows: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """解析 Parquet 表并在物化行数据前检查元数据行数，避免压缩文件放大耗尽内存。"""
    try:
        parquet_file = arrow_parquet.ParquetFile(pa.BufferReader(content))
        row_count = parquet_file.metadata.num_rows
        if row_count > max_rows:
            raise DomainError(f"文件超过单次解析允许的 {max_rows} 行上限")
        table = parquet_file.read()
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError("Parquet 文件无法解析或结构已损坏") from exc
    rows = [_arrow_row(row) for row in table.to_pylist()]
    return rows, {
        "row_group_count": parquet_file.metadata.num_row_groups,
        "column_count": table.num_columns,
    }


def _read_orc(content: bytes, max_rows: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """解析 ORC 表并在读取完整内容前检查行数，保持与其他受管文件相同的容量边界。"""
    try:
        orc_file = pyorc.Reader(io.BytesIO(content))
        if len(orc_file) > max_rows:
            raise DomainError(f"文件超过单次解析允许的 {max_rows} 行上限")
        schema = cast(Any, orc_file.schema)
        field_names = list(schema.fields)
        rows: list[dict[str, Any]] = []
        for raw_values in orc_file:
            values = cast(tuple[Any, ...], raw_values)
            rows.append(_arrow_row(dict(zip(field_names, values, strict=True))))
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError("ORC 文件无法解析或结构已损坏") from exc
    return rows, {
        "stripe_count": orc_file.num_of_stripes,
        "column_count": len(field_names),
    }


def _read_xml(content: bytes, max_rows: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """把 XML 根节点下的重复记录元素展开为行，并禁止实体扩展等不安全 XML 特性。"""
    try:
        root = SafeElementTree.fromstring(content)
    except Exception as exc:
        raise DomainError("XML 文件无法解析或包含不安全结构") from exc
    elements = list(root)
    if not elements:
        raise DomainError("XML 根节点下没有可解析的记录元素")
    rows: list[dict[str, Any]] = []
    for element in elements:
        row = _flatten_xml_element(element)
        if row:
            rows.append(row)
            _require_row_capacity(rows, max_rows)
    return rows, {
        "root_element": _xml_local_name(root.tag),
        "record_element": _xml_local_name(elements[0].tag),
    }


def _arrow_row(row: dict[str, Any]) -> dict[str, Any]:
    """把 Arrow 可能返回的日期、Decimal、二进制和嵌套值转换为统一 JSON 行标量。"""
    return {str(key): _json_value(value) for key, value in row.items()}


def _flatten_xml_element(element: Any) -> dict[str, Any]:
    """保守展开单条 XML 记录；重复或深层节点以规范 JSON 保存，避免静默丢字段。"""
    row: dict[str, Any] = {
        f"attr_{_xml_local_name(key)}": value for key, value in element.attrib.items()
    }
    children = list(element)
    if not children:
        text = (element.text or "").strip()
        return {"content": text} if text else row
    grouped: dict[str, list[Any]] = {}
    for child in children:
        grouped.setdefault(_xml_local_name(child.tag), []).append(child)
    for name, values in grouped.items():
        converted = [_xml_element_value(value) for value in values]
        row[name] = converted[0] if len(converted) == 1 else converted
    return row


def _xml_element_value(element: Any) -> Any:
    """将 XML 子节点转换为标量或嵌套对象，使重复节点信息完整保留。"""
    children = list(element)
    if not children and not element.attrib:
        return _parse_text_scalar((element.text or "").strip())
    nested = _flatten_xml_element(element)
    text = (element.text or "").strip()
    if text:
        nested["content"] = text
    return nested


def _xml_local_name(tag: str) -> str:
    """移除 XML 命名空间前缀，避免同一业务字段因部署命名空间不同而产生不稳定列名。"""
    return tag.rsplit("}", 1)[-1]


def _read_markdown(content: bytes, max_rows: int) -> list[dict[str, Any]]:
    """把 Markdown 标题、段落和列表项转换为稳定 DocumentElement 行。"""
    rows: list[dict[str, Any]] = []
    for raw_line in _decode_text(content).splitlines():
        line = raw_line.strip()
        if not line or line.startswith("```"):
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            element_type = f"heading_{min(level, 6)}"
            text = line[level:].strip()
        elif re.match(r"^[-*+]\s+", line):
            element_type = "list_item"
            text = re.sub(r"^[-*+]\s+", "", line)
        else:
            element_type = "paragraph"
            text = line
        if text:
            rows.append(
                {"element_index": len(rows) + 1, "element_type": element_type, "content": text}
            )
            _require_row_capacity(rows, max_rows)
    return rows


def _read_pdf(content: bytes, max_rows: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """提取文本型 PDF 的逐页内容；扫描件无文本时明确失败，不伪装 OCR 能力。"""
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise DomainError("暂不支持加密 PDF")
        if len(reader.pages) > max_rows:
            raise DomainError(f"文件超过单次解析允许的 {max_rows} 行上限")
        rows = [
            {"page_number": index, "content": (page.extract_text() or "").strip()}
            for index, page in enumerate(reader.pages, start=1)
        ]
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError("PDF 文件无法解析或结构已损坏") from exc
    rows = [row for row in rows if row["content"]]
    if not rows:
        raise DomainError("PDF 未提取到文本；扫描件和 OCR 暂不支持")
    return rows, {"page_count": len(reader.pages), "document_model": "pdf_pages"}


def _require_row_capacity(rows: Sequence[Any], max_rows: int) -> None:
    """在解析循环内尽早拒绝超限文件，避免小体积高行数内容耗尽 Worker 内存。"""
    if len(rows) > max_rows:
        raise DomainError(f"文件超过单次解析允许的 {max_rows} 行上限")


def _normalize_rows(
    rows: Sequence[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    """稳定规范化字段名和值，使任意用户表头都能安全进入 JSONL 与 HOCON Schema。"""
    originals: list[str] = []
    for row in rows:
        for key in row:
            name = str(key)
            if name not in originals:
                originals.append(name)
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for index, original in enumerate(originals, start=1):
        base = _UNSAFE_COLUMN.sub("_", original.strip()).strip("_").lower()
        if not base or base[0].isdigit():
            base = f"field_{index}_{base}".rstrip("_")
        candidate = base
        suffix = 2
        while candidate in used:
            candidate = f"{base}_{suffix}"
            suffix += 1
        mapping[original] = candidate
        used.add(candidate)
    normalized = [
        {
            mapping[str(key)]: (
                str(value)
                if value is not None and SENSITIVE_COLUMN.search(str(key))
                else _json_value(value)
            )
            for key, value in row.items()
        }
        for row in rows
    ]
    return normalized, mapping


def _json_value(value: Any) -> Any:
    """把解析器值约束为 JSON 标量；嵌套对象编码成规范字符串以保持行模型扁平。"""
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)
    return str(value)


def _parse_text_scalar(value: str | None) -> Any:
    """为 CSV 文本执行保守标量推断，空值、整数、浮点和布尔之外保持字符串。"""
    if value is None or value == "":
        return None
    normalized = value.strip()
    if normalized.casefold() in {"true", "false"}:
        return normalized.casefold() == "true"
    if re.fullmatch(r"[-+]?\d+", normalized):
        try:
            return int(normalized)
        except ValueError:
            return value
    if re.fullmatch(r"[-+]?(?:\d+\.\d*|\d*\.\d+)(?:[eE][-+]?\d+)?", normalized):
        try:
            return float(normalized)
        except ValueError:
            return value
    return value


def _infer_schema(rows: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    """按全部规范化行推导保守 SeaTunnel 类型，混合类型统一提升为 string。"""
    columns: list[str] = []
    for row in rows:
        for column in row:
            if column not in columns:
                columns.append(column)
    schema: list[dict[str, Any]] = []
    for column in columns:
        values = [row.get(column) for row in rows]
        kinds = {_value_kind(value) for value in values if value is not None}
        if not kinds or "string" in kinds or (len(kinds) > 1 and kinds != {"bigint", "double"}):
            column_type = "string"
        elif "double" in kinds:
            column_type = "double"
        else:
            column_type = next(iter(kinds))
        schema.append(
            {
                "name": column,
                "type": column_type,
                "nullable": any(value is None for value in values),
            }
        )
    return schema


def _value_kind(value: Any) -> str:
    """把 JSON 标量映射为 SeaTunnel 文件 Schema 支持的基础类型。"""
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "bigint"
    if isinstance(value, float):
        return "double"
    return "string"
